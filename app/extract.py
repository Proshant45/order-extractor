"""File -> structured order. Memory-conscious: page by page, bounded image sizes."""
from __future__ import annotations
import io, os, re, zipfile
from pathlib import Path
from pydantic import ValidationError
from . import config, llm, parsers
from .schema import Order, warnings_for

class Skip(Exception):
    """File is not processable; recorded as skipped with a reason, never fatal."""
    def __init__(self, msg, need_ai=False):
        super().__init__(msg); self.need_ai = need_ai

def sniff(path: Path) -> str:
    with open(path, "rb") as f:
        head = f.read(8)
    ext = path.suffix.lower()
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\x89PNG") or head[:3] == b"\xff\xd8\xff" or head[:4] in (b"GIF8", b"II*\x00", b"MM\x00*") or head[:4] == b"RIFF":
        return "image"
    if head.startswith(b"PK"):
        if ext == ".xlsx": return "xlsx"
        if ext == ".docx": return "docx"
        try:
            with zipfile.ZipFile(path) as z:
                names = z.namelist()[:50]
        except zipfile.BadZipFile:
            raise Skip("corrupt zip container")
        if any(n.startswith("xl/") for n in names): return "xlsx"
        if any(n.startswith("word/") for n in names): return "docx"
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls" if ext in (".xls", "") else "unknown"
    if ext in (".csv", ".tsv"):
        return "csv"
    if ext == ".eml":
        return "eml"
    if ext in (".txt", ".md", ".json", ".xml", ".html", ".htm"):
        return "text"
    if ext in (".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp", ".heic"):
        return "image"  # extension says image even if magic is odd: let AI reader try
    return "unknown"

def _png_bytes_from_image(path: Path) -> bytes:
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = 60_000_000  # decompression-bomb guard
    try:
        with Image.open(path) as im:
            im.draft("RGB", (config.MAX_IMAGE_PX, config.MAX_IMAGE_PX))  # cheap JPEG downscale on decode
            im = im.convert("RGB")
            im.thumbnail((config.MAX_IMAGE_PX, config.MAX_IMAGE_PX))
            buf = io.BytesIO(); im.save(buf, "PNG", optimize=True)
            return buf.getvalue()
    except Exception as e:
        raise Skip(f"unreadable image: {e.__class__.__name__}")

def _need_llm(log, what):
    if config.llm_provider() is None:
        raise Skip(f"{what} needs the AI reader, but the AI worker queue is not enabled on this server", need_ai=True)
    log(f"AI reader ({config.llm_provider()}) for {what}")

def extract_pdf(path: Path, log, ctx=None) -> tuple[dict, str]:
    import pdfplumber, pymupdf
    texts, scanned = [], []
    try:
        pdf = pdfplumber.open(path)
    except Exception as e:
        raise Skip(f"corrupt PDF: {e.__class__.__name__}")
    with pdf:
        n = len(pdf.pages)
        if n > config.MAX_PAGES:
            log(f"PDF has {n} pages, only first {config.MAX_PAGES} processed", "warn")
        for i, page in enumerate(pdf.pages[:config.MAX_PAGES]):
            t = page.extract_text(layout=True) or ""
            page.flush_cache()           # drop parsed objects for this page
            texts.append(t)
            if len(t.strip()) < 30:
                scanned.append(i)
    if ctx is not None:
        ctx["category"] = "pdf_scanned" if len(scanned) > len(texts) / 2 else "pdf_text"
    full = "\n\f".join(texts)
    log(f"PDF pages={len(texts)} text_pages={len(texts)-len(scanned)} scanned_pages={len(scanned)}")
    if ctx and ctx.get("force_ai"):
        return _forced_ai_pdf(path, texts, scanned, log, ctx)
    if parsers.looks_like_sap_po(full):
        log("deterministic parser: sap_po")
        return parsers.parse_sap_po(full), "pdf_text:sap_po"
    _need_llm(log, "PDF with unknown layout or scanned pages")
    pages = []
    doc = pymupdf.open(path)
    try:
        for i in range(min(len(doc), config.MAX_PAGES)):
            if i in scanned:
                pix = doc[i].get_pixmap(dpi=config.RENDER_DPI)   # one page image in memory at a time
                png = pix.tobytes("png"); del pix
                pages.append(llm.read_page(image_png=png)); del png
            else:
                pages.append(llm.read_page(text=texts[i]))
            log(f"AI read page {i+1}")
    finally:
        doc.close()
    return llm.merge_pages(pages), "pdf_llm"

def _guarded(read, text, log, ctx):
    from . import guard
    d, errs, n = guard.ai_extract_with_guard(read, text, log)
    ctx["guard_errors"], ctx["attempts"] = errs, n
    return d

def _forced_ai_pdf(path, texts, scanned, log, ctx):
    import pymupdf
    _need_llm(log, "forced AI read of PDF")
    if not scanned:
        full = "\n\f".join(texts)
        return _guarded(lambda extra: llm.read_page(text=full, extra=extra), full, log, ctx), "pdf_ai_text"
    pages = []
    doc = pymupdf.open(path)
    try:
        for i in range(min(len(doc), config.MAX_PAGES)):
            if i in scanned:
                png = doc[i].get_pixmap(dpi=config.RENDER_DPI).tobytes("png")
                pages.append(_guarded(lambda extra, png=png: llm.read_page(image_png=png, extra=extra), None, log, ctx))
            else:
                pages.append(_guarded(lambda extra, t=texts[i]: llm.read_page(text=t, extra=extra), texts[i], log, ctx))
            log(f"AI read page {i+1}")
    finally:
        doc.close()
    return llm.merge_pages(pages), "pdf_ai_vision"

def _as_text(path: Path, kind: str) -> str | None:
    if kind == "csv":
        return path.read_text(errors="replace")[:200000]
    if kind == "xlsx":
        return "\n".join("\t".join("" if c is None else str(c) for c in r) for _, r in zip(range(5000), parsers.iter_xlsx(path)))
    if kind == "docx":
        import docx
        d = docx.Document(str(path))
        return "\n".join([p.text for p in d.paragraphs] + ["\t".join(c.text for c in row.cells) for t in d.tables for row in t.rows])
    if kind == "xls":
        import xlrd
        wb = xlrd.open_workbook(str(path), on_demand=True)
        sh = wb.sheet_by_index(0)
        return "\n".join("\t".join(str(c.value) for c in sh.row(i)) for i in range(min(sh.nrows, 5000)))
    if kind == "eml":
        import email
        from email import policy
        m = email.message_from_binary_file(open(path, "rb"), policy=policy.default)
        b = m.get_body(preferencelist=("plain", "html"))
        body = b.get_content() if b else ""
        if b is not None and b.get_content_type() == "text/html":
            body = re.sub(r"<[^>]+>", " ", body)
        return f"Subject: {m['subject']}\nFrom: {m['from']}\n\n{body}"[:200000]
    if kind in ("text", "unknown"):
        raw = path.read_bytes()[:400000]
        if b"\x00" in raw[:4096]:
            return None
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            return raw.decode("latin-1") if sum(c < 9 for c in raw[:4096]) == 0 else None
    return None

def ai_primary(ctx) -> bool:
    import os
    if ctx.get("force_ai"):
        return True
    return os.getenv("PRIMARY_EXTRACTOR", "ai").lower() == "ai" and config.llm_provider() is not None

PO_HDR = re.compile(r"(?:Purchase\s*Order|P\.?\s*O\.?)\s*(?:No\.?|Number|#)?\s*[:#]\s*([A-Z]*\d[\w\-/]{3,})", re.I)

def _po_groups(texts: list[str]) -> list[tuple[int, int, int]]:
    """[(first_page, end_page, n_pos)]. One group unless the doc has more than AI_SPLIT_MIN_POS distinct PO headers;
    then pages are cut only where a new PO starts, AI_POS_PER_READ POs per read."""
    first = {}
    for i, t in enumerate(texts):
        for m in PO_HDR.finditer(t or ""):
            first.setdefault(m.group(1).upper(), i)
    if len(first) <= int(os.getenv("AI_SPLIT_MIN_POS", "6")):
        return [(0, len(texts), len(first))]
    starts = sorted(set(first.values())); starts[0] = 0
    per = int(os.getenv("AI_POS_PER_READ", "4"))
    cuts = starts[::per] + [len(texts)]
    return [(cuts[k], cuts[k + 1], len([v for v in first.values() if cuts[k] <= v < cuts[k + 1]]) or per) for k in range(len(cuts) - 1)]

def _pdf_units(path: Path, log, ctx):
    import pdfplumber, pymupdf
    try:
        pdf = pdfplumber.open(path)
    except Exception as e:
        raise Skip(f"corrupt PDF: {e.__class__.__name__}")
    texts = []
    with pdf:
        for page in pdf.pages[:config.MAX_PAGES]:
            texts.append(page.extract_text(layout=True) or ""); page.flush_cache()
    scanned = [i for i, t in enumerate(texts) if len(t.strip()) < 30]
    ctx["category"] = "pdf_scanned" if len(scanned) > len(texts) / 2 else "pdf_text"
    log(f"PDF pages={len(texts)} scanned_pages={len(scanned)}")
    if not scanned:   # whole document in one read; split only multi-PO docs, on PO boundaries
        full = "\n\f".join(texts)
        det = parsers.parse_sap_po(full) if parsers.looks_like_sap_po(full) else None
        groups = _po_groups(texts)
        if len(groups) <= 1:
            yield {"text": full, "page": "1-%d" % len(texts), "det": det}
            return
        log(f"{sum(g[2] for g in groups)} PO headers; reading in {len(groups)} PO-boundary groups")
        for a0, a1, _ in groups:
            yield {"text": "\n\f".join(texts[a0:a1]), "page": "%d-%d" % (a0 + 1, a1), "det": None}
        return
    doc = pymupdf.open(path)
    try:
        for i in range(min(len(doc), config.MAX_PAGES)):
            if i in scanned:
                pix = doc[i].get_pixmap(dpi=config.RENDER_DPI); png = pix.tobytes("png"); del pix
                yield {"image": png, "page": str(i + 1), "det": None}
            else:
                yield {"text": texts[i], "page": str(i + 1), "det": None}
    finally:
        doc.close()

def extract_ai_primary(path: Path, kind: str, log, ctx) -> tuple[dict, str]:
    from . import consensus
    _need_llm(log, f"{kind} (AI-primary dual read)")
    if kind == "pdf":
        units = _pdf_units(path, log, ctx)
    elif kind in ("image", "unknown") and not (kind == "unknown" and _as_text(path, kind)):
        units = [{"image": _png_bytes_from_image(path), "page": "1", "det": None}]
    else:
        try:
            txt = _as_text(path, kind)
        except Exception as e:
            raise Skip(f"corrupt {kind}: {e.__class__.__name__}")
        if not txt or not txt.strip():
            raise Skip(f"{kind}: no readable text")
        det = None
        try:
            if kind in ("csv", "xlsx"):
                det = parsers.parse_table_rows(parsers.iter_csv(path) if kind == "csv" else parsers.iter_xlsx(path))
        except Exception:
            det = None
        units = [{"text": txt, "page": "1", "det": det}]
    A, B, dis, arith, gerr, status, used_det = [], [], [], [], [], "accepted", False
    owner = getattr(llm.CURRENT, "owner", None)
    def _one(u):
        llm.CURRENT.owner = owner
        return consensus.dual_read(text=u.get("text"), image_png=u.get("image"), det=u["det"], log=log)
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=int(os.getenv("AI_PAGE_PARALLEL", "2"))) as ex:
        units = list(units)
        results = list(ex.map(_one, units))      # pages read concurrently, merged in page order
    for u, r in zip(units, results):
        used_det = used_det or u["det"] is not None
        A.append(r["votes"]["a"]["data"]); B.append(r["votes"]["b"]["data"] or {})
        dis += [dict(x, page=u["page"]) for x in r["disagreements"]]
        arith += r["arithmetic_errors"]
        gerr += [f"[{v}] {e}" for v in ("a", "b") for e in r["votes"][v]["guard_errors"]]
        ctx.setdefault("models", {"a": r["votes"]["a"]["model"], "b": r["votes"]["b"]["model"]})
        ctx["attempts"] = max(ctx.get("attempts") or 0, r["votes"]["a"]["attempts"], r["votes"]["b"]["attempts"])
        if r["status"] != "accepted":
            status = "needs_review"
    a, b = llm.merge_pages(A), llm.merge_pages(B)
    if not (a.get("pos") or []):
        status = "needs_review"; arith.append("no PO / quantity data extracted (never auto-accepted)")
    ctx["consensus"] = {"status": status, "disagreements": dis, "arithmetic_errors": arith, "guard_errors": gerr,
                        "models": ctx.get("models"), "version_a": a, "version_b": b,
                        "deterministic_vote": used_det}
    ctx["guard_errors"] = arith + gerr
    return a, f"{kind}_ai_dual"

def preview_text(path: Path, kind: str) -> str | None:
    """Cheap, bounded text for classification (first 3 pages / first rows)."""
    try:
        if kind == "pdf":
            import pdfplumber
            out = []
            with pdfplumber.open(path) as pdf:
                for pg in pdf.pages[:3]:
                    out.append(pg.extract_text() or ""); pg.flush_cache()
            return "\n".join(out)
        if kind == "xlsx":
            return "\n".join("\t".join("" if c is None else str(c) for c in r) for _, r in zip(range(150), parsers.iter_xlsx(path)))
        return (_as_text(path, kind) or "")[:20000]
    except Exception:
        return None

class NotOrder(Exception):
    """Document classified as non-order (tech pack, BOM, other): stored with type + refs, not forced into the order schema."""

def extract_file(path: Path, kind: str, log, ctx=None) -> tuple[dict, str]:
    ctx = ctx if ctx is not None else {}
    from . import classify
    pv = preview_text(path, kind)
    dt, scores = classify.classify(ctx.get("name", path.name), pv)
    ctx["doc_type"] = dt; ctx["refs"] = classify.refs(pv)
    log(f"doc type: {dt} {dict((k, v) for k, v in scores.items() if v)}")
    if dt not in classify.ORDER_TYPES and not ctx.get("force_order") and kind != "image" and (pv or "").strip():
        raise NotOrder(dt)
    if ai_primary(ctx):
        return extract_ai_primary(path, kind, log, ctx)
    if kind in ("xls", "eml", "text") or (kind == "unknown" and _as_text(path, kind)):
        txt = _as_text(path, kind)  # readable-but-unsupported formats go to the AI reader
        if not txt or not txt.strip():
            raise Skip(f"{kind}: no readable text")
        _need_llm(log, f"{kind} file")
        return _guarded(lambda extra: llm.read_page(text=txt, extra=extra), txt, log, ctx), f"{kind}_ai"
    if ctx.get("force_ai") and kind in ("csv", "xlsx", "docx"):
        try:
            txt = _as_text(path, kind)
        except Exception as e:
            raise Skip(f"corrupt {kind}: {e.__class__.__name__}")
        _need_llm(log, f"forced AI read of {kind}")
        return _guarded(lambda extra: llm.read_page(text=txt, extra=extra), txt, log, ctx), f"{kind}_ai"
    if ctx.get("force_ai") and kind == "image":
        png = _png_bytes_from_image(path)
        _need_llm(log, "image")
        return _guarded(lambda extra: llm.read_page(image_png=png, extra=extra), None, log, ctx), "image_ai"
    if kind == "pdf":
        return extract_pdf(path, log, ctx)
    if kind in ("csv", "xlsx"):
        try:
            rows = parsers.iter_csv(path) if kind == "csv" else parsers.iter_xlsx(path)
            d = parsers.parse_table_rows(rows)
        except Exception as e:
            raise Skip(f"corrupt {kind}: {e.__class__.__name__}")
        if d:
            return d, f"{kind}:table"
        _need_llm(log, f"{kind} without a Color/size header")
        rows = parsers.iter_csv(path) if kind == "csv" else parsers.iter_xlsx(path)
        txt = "\n".join("\t".join("" if c is None else str(c) for c in r) for _, r in zip(range(2000), rows))
        return llm.merge_pages([llm.read_page(text=txt)]), f"{kind}_llm"
    if kind == "docx":
        import docx
        try:
            d = docx.Document(str(path))
        except Exception as e:
            raise Skip(f"corrupt docx: {e.__class__.__name__}")
        paras = [p.text for p in d.paragraphs]
        meta_rows = [[p.split(":", 1)[0] + ":", p.split(":", 1)[1].strip()] for p in paras if ":" in p]
        for t in d.tables:
            r = parsers.parse_table_rows(meta_rows + [[c.text for c in row.cells] for row in t.rows])
            if r:
                return r, "docx:table"
        _need_llm(log, "docx")
        txt = "\n".join(paras + ["\t".join(c.text for c in row.cells) for t in d.tables for row in t.rows])
        return llm.merge_pages([llm.read_page(text=txt)]), "docx_llm"
    if kind in ("image", "unknown"):
        png = _png_bytes_from_image(path)  # unknown files: try as image; Skip if not decodable
        _need_llm(log, "image")
        return _guarded(lambda extra: llm.read_page(image_png=png, extra=extra), None, log, ctx), "image_ai"
    raise Skip(f"unsupported type {kind}")

def validate(data: dict) -> tuple[Order | None, list[str], list[str]]:
    data = {k: v for k, v in data.items() if not k.startswith("_")}
    try:
        o = Order.model_validate(data)
    except ValidationError as e:
        return None, [], [f"{'.'.join(map(str, er['loc']))}: {er['msg']}" for er in e.errors()]
    return o, warnings_for(o), []
