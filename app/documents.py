import hashlib
import json
import logging
import os
from pathlib import Path
import secrets
import shutil
import zipfile
from . import config, db, extract, zipio


def retain_source(path,sha=None):
    path=Path(path)
    if sha is None:
        digest=hashlib.sha256()
        with path.open('rb') as src:
            for chunk in iter(lambda:src.read(config.CHUNK),b''):digest.update(chunk)
        sha=digest.hexdigest()
    directory=config.DATA_DIR/'documents';directory.mkdir(parents=True,exist_ok=True)
    target=directory/(sha+path.suffix.lower()[:10])
    if not target.exists():
        tmp=directory/('.'+sha+'-'+secrets.token_hex(4))
        try:
            shutil.copyfile(path,tmp)
            with tmp.open('rb') as f:os.fsync(f.fileno())
            tmp.replace(target)
        finally:tmp.unlink(missing_ok=True)
    return target


def attach_source(fid,path,sha):
    target=retain_source(path,sha)
    with db.conn() as c:
        c.execute('UPDATE files SET source_path=? WHERE id=? OR sha256=?',(str(target),fid,sha))
    return target


def source_for_file(file):
    if not file.get('source_path'):return None
    path=Path(file['source_path'])
    if not path.is_file() or not path.resolve().is_relative_to((config.DATA_DIR/'documents').resolve()):return None
    return path


def restore_retained_batch_sources():
    with db.conn() as c:rows=[dict(row) for row in c.execute('SELECT name,input_path FROM batches WHERE input_path IS NOT NULL')]
    for batch in rows:
        path=Path(batch['input_path'])
        if not path.is_file():continue
        try:
            digest=hashlib.sha256()
            with path.open('rb') as src:
                for chunk in iter(lambda:src.read(config.CHUNK),b''):digest.update(chunk)
            sha=digest.hexdigest()
            with db.conn() as c:
                row=c.execute('SELECT id FROM files WHERE name=? AND sha256=? ORDER BY id DESC LIMIT 1',(batch['name'],sha)).fetchone()
            kind=extract.sniff(path)
            fid=row['id'] if row else db.add_file(batch['name'],kind=kind,size=path.stat().st_size,sha=sha,status='pending',reason='Extraction has not run for this file')
            attach_source(fid,path,sha)
            if kind=='zip':index_archive(path,fid)
        except Exception:logging.exception('Could not restore source documents for %s',batch['name'])


def index_archive(path,parent_id):
    for name,tmp,reason in zipio.iter_zip(path):
        if zipio.junk_reason(name):
            if tmp:tmp.unlink(missing_ok=True)
            continue
        try:
            with db.conn() as c:
                row=c.execute('SELECT id FROM files WHERE parent_id=? AND name=? ORDER BY id DESC LIMIT 1',(parent_id,name)).fetchone()
            if tmp:
                target=retain_source(tmp);sha=target.name.split('.')[0]
                try:kind=extract.sniff(tmp)
                except extract.Skip:kind='unknown'
                fid=row['id'] if row else db.add_file(name,kind=kind,size=tmp.stat().st_size,sha=sha,status='pending',parent_id=parent_id,reason='Extraction has not run for this file')
                attach_source(fid,tmp,sha)
            elif not row:
                db.add_file(name,status='skipped',parent_id=parent_id,reason=reason)
        finally:
            if tmp:tmp.unlink(missing_ok=True)


DOC_TYPES = {"purchase_order":"Purchase order", "sales_contract":"Sales contract", "proforma_invoice":"Proforma invoice", "tech_pack":"Tech pack", "bom":"Bill of materials", "control_sheet":"Control sheet", "other":"Other document"}


def describe_file(file):
    status=file["status"]; supporting=file.get("doc_type") in ("tech_pack","bom","other")
    labels={"done":"Order ready", "cached":"Already extracted" if file.get("order_id") else "Supporting document", "classified":"Supporting document", "pending":"Not extracted", "processing":"Processing", "queued_ai":"Waiting for retry" if file.get('reason') else "Waiting for worker", "failed":"Extraction failed", "skipped":"Not processed", "cancelled":"Cancelled", "retried":"Retried", "edited":"Order edited"}
    label=labels.get(status,status.replace("_"," ").title())
    if status=="skipped" and "processing limit" in (file.get("reason") or "").lower():label="Limit reached"
    if status in ("classified","cached") and not file.get("order_id"):
        explanation="Identified as " + DOC_TYPES.get(file.get("doc_type"),"a supporting document") + ". It is available to inspect; no order was created."
    elif status=="cached" and file.get("order_id"):
        explanation="An identical copy was extracted earlier. Open the existing order or inspect the source document."
    elif status=="processing":explanation="The document is being read and validated. Follow its progress on the Order Book."
    elif status=="queued_ai":explanation=(file.get('reason') or 'Waiting for the AI worker.')+' No completed order is available yet.'
    elif status=="done":explanation="An order was created. Open it to check quantities, warnings and missing fields."
    else:explanation=file.get("reason") or "Inspect the document and diagnostic details below."
    return {**file,"label":label,"explanation":explanation,"type_label":DOC_TYPES.get(file.get("doc_type"),"Not identified yet"),"source_available":source_for_file(file) is not None,"display_name":Path(file["name"]).name,"folder":str(Path(file["name"]).parent) if "/" in file["name"] else ""}


def preview(file,path):
    kind=file.get("kind")
    if kind=="zip":
        with db.conn() as c:
            related={r["name"]:r["id"] for r in c.execute("SELECT id,name FROM files WHERE parent_id=? ORDER BY id",(file["id"],))}
        with zipfile.ZipFile(path) as archive:
            entries=[{"name":info.filename,"size":info.file_size,"file_id":related.get(info.filename)} for info in archive.infolist() if not info.is_dir() and not zipio.junk_reason(info.filename)]
        return {"kind":"archive","entries":entries}
    if kind in ("pdf","image"):return {"kind":kind}
    if kind=="xlsx":
        import openpyxl
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        try:
            sheets=[]
            for ws in wb.worksheets:
                rows=[]
                for row in ws.iter_rows(max_row=min(ws.max_row or 100,100),max_col=min(ws.max_column or 40,40),values_only=True):
                    if any(value is not None for value in row):rows.append(["" if value is None else str(value) for value in row])
                sheets.append({"name":ws.title,"rows":rows,"truncated":(ws.max_row or 0)>100 or (ws.max_column or 0)>40})
            return {"kind":"sheets","sheets":sheets}
        finally:wb.close()
    if kind in ("csv","text","unknown"):
        with path.open('r',errors='replace') as source:text=source.read(20001)
    else:text=extract._as_text(path,kind) or ''
    return {"kind":"text","text":text[:20000],"truncated":len(text)>20000}


def library_files():
    out=[];seen=set()
    for file in db.list_files(-1):
        key=(file["name"],file.get("sha256") or file["id"])
        if key in seen:continue
        seen.add(key);out.append(describe_file(file))
    return out


def stored_data(file):
    order=db.get_order(file['order_id']) if file.get('order_id') else None
    with db.conn() as c:
        raw=c.execute('SELECT raw_json FROM orders WHERE id=?',(file['order_id'],)).fetchone() if order else None
    return {'file_id':file['id'],'status':file['status'],'reason':file.get('reason'),
            'metadata':json.loads(file.get('meta') or '{}'),'order':order,
            'original_extraction':json.loads(raw['raw_json'] or 'null') if raw else None}
