import json
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from app import db, main, config, llm, consensus, guard
from app.schema import Order, POLine

@pytest.fixture
def client():
    db.init()
    with TestClient(main.app, base_url='https://testserver') as c:
        assert c.post('/login', data={'username':'tester','password':'pw-123456'}, follow_redirects=False).status_code == 303
        yield c


def test_draft_and_saved_survive_readback(client):
    body = {'buyer_name':'Draft buyer','status':'saved','pos':[]}
    oid = client.post('/api/orders', json=body).json()['id']
    assert client.get(f'/api/orders/{oid}').json()['status'] == 'saved'
    body['status'] = 'draft'
    assert client.put(f'/api/orders/{oid}', json=body).json()['ok']
    assert client.get(f'/api/orders/{oid}').json()['status'] == 'draft'
    assert client.put(f'/api/orders/{oid}', json={**body,'status':'nonsense'}).status_code == 422


@pytest.mark.parametrize('q', [-1, 1.5, '1.5', 'nan', 'inf'])
def test_no_silent_rounding_of_quantities(q):
    with pytest.raises(ValueError):
        POLine.model_validate({'color':'Navy','order_qty':q})
    with pytest.raises(ValueError):
        Order.model_validate({'pos':[{'po_number':'PO','stated_total_qty':q}]})


def test_raw_dates_and_image_survive_save(client):
    body = {'style_image':'data:image/png;base64,aGVsbG8=', 'pos':[{'po_number':'RAW', 'ship_date':'13/09','lines':[]}]}
    oid = client.post('/api/orders', json=body).json()['id']
    o = client.get(f'/api/orders/{oid}').json()
    assert o['style_image'] == body['style_image']
    assert o['pos'][0]['date_raw'] == {'ship_date':'13/09'}


def test_expired_lease_cannot_renew_or_finish():
    db.init()
    jid = db.enqueue_job('prompt','model','text',None,owner='tester')
    job = db.claim_job('box','tester')
    # Existing tests may leave other queued jobs; claim our job explicitly for this check.
    while job and job['id'] != jid:
        db.finish_job(job['id'],job['lease_id'], result='{}')
        job = db.claim_job('box','tester')
    with db.conn() as c:
        c.execute('UPDATE ai_jobs SET lease_until=? WHERE id=?',(time.time()-1,jid))
    assert not db.renew_job(jid, job['lease_id'])
    assert db.finish_job(jid,job['lease_id'],result='{}') == 'stale_lease'
    new = db.claim_job('box','tester')
    assert new['id'] == jid and new['lease_id'] != job['lease_id']
    assert db.finish_job(jid,new['lease_id'],result='{}') == 'ok'


def test_comparison_does_not_hide_different_styles():
    a = {'pos':[{'po_number':'1','lines':[{'style':'A','color':'Navy','sizes':{'S':10}}, {'style':'B','color':'Navy','sizes':{'S':20}}]}]}
    b = json.loads(json.dumps(a)); b['pos'][0]['lines'][0]['sizes']['S'] = 999
    assert consensus.compare(a,b)


def test_missing_size_breakdown_uses_stated_line_quantity():
    d = {'pos':[{'po_number':'1','stated_total_qty':1200,'stated_total_value':3000,
                'lines':[{'color':'Navy','sizes':{},'order_qty':1200,'rate':2.5,'net_value':3000}]}]}
    assert guard.arithmetic_errors(d) == []


def test_reader_refuses_silent_text_truncation(monkeypatch):
    monkeypatch.setattr(config,'llm_provider',lambda:'queue')
    monkeypatch.setattr(llm,'queue_read',lambda *a, **k:'{"pos":[]}')
    with pytest.raises(ValueError, match='too large'):
        llm.read_page(text='x'*60001)


def test_api_token_jobs_belong_to_app_account(client, monkeypatch):
    monkeypatch.setenv('API_TOKEN','testing-api-token')
    owners=[]
    monkeypatch.setattr(main,'_run_upload',lambda *a,**kw: _capture_owner(owners,{'owner':a[3]}))
    client.cookies.clear()
    assert client.post('/api/upload?wait=0',headers={'Authorization':'Bearer testing-api-token'},files={'file':('x.txt',b'x')}).status_code == 200
    assert owners == ['tester']

async def _capture_owner(owners,kw):
    owners.append(kw.get('owner'))
    return {'batch_id':'test','status':'processing'}


def test_interrupted_batch_replays_retained_input(client, monkeypatch, tmp_path):
    source = tmp_path/'retained.csv'; source.write_text('Buyer:,Recovery\nPO:,RECOVERY\nColor,S,M\nNavy,10,20\n')
    bid = 'restart-test'
    db.save_batch(bid,'processing','retained.csv')
    with db.conn() as c:
        c.execute('UPDATE batches SET input_path=?, input_size=?, input_sha=?, owner=? WHERE id=?',
                  (str(source),source.stat().st_size,main._sha256(source),'tester',bid))
    main._recover_batches()
    deadline = time.monotonic()+10
    while time.monotonic()<deadline and db.get_batch(bid)['status']=='processing':
        time.sleep(.05)
    b = db.get_batch(bid)
    assert b['status']=='done' and b['result']['items'][0]['status']=='done'
    assert db.get_order(b['result']['items'][0]['order_id'])['buyer_name']=='Recovery'
    assert not source.exists()


def test_partial_opencode_text_is_not_a_completed_answer(monkeypatch, tmp_path):
    from app import opencode_provider as oc
    exe = tmp_path/'fake-opencode'
    exe.write_text('#!/usr/bin/env python3\nimport json\nprint(json.dumps({"type":"text","part":{"text":"{\\\"pos\\\":[]}"}}),flush=True)\n')
    exe.chmod(0o755)
    monkeypatch.setattr(oc,'binary',lambda:str(exe))
    with pytest.raises(RuntimeError,match='completion'):
        oc.run('prompt',timeout=2)


def test_opencode_partial_line_cannot_bypass_deadline(monkeypatch, tmp_path):
    from app import opencode_provider as oc
    exe = tmp_path/'fake-opencode'
    exe.write_text('#!/usr/bin/env python3\nimport sys,time\nsys.stdout.write("{partial");sys.stdout.flush();time.sleep(2)\n')
    exe.chmod(0o755)
    monkeypatch.setattr(oc,'binary',lambda:str(exe))
    start=time.monotonic()
    with pytest.raises(TimeoutError):
        oc.run('prompt',timeout=.2)
    assert time.monotonic()-start < 1.5


def test_worker_keeps_result_when_server_rejects_post(monkeypatch, tmp_path):
    import httpx
    monkeypatch.setenv('APP_URL','http://local.test')
    monkeypatch.setenv('WORKER_TOKEN','x'*40)
    from worker import box_worker as worker
    monkeypatch.setattr(worker,'OUTBOX',tmp_path)
    monkeypatch.setattr(worker.time,'sleep',lambda _:None)
    body={'id':42,'lease_id':'abc','text':'{"pos":[]}'}
    with httpx.Client(transport=httpx.MockTransport(lambda req: httpx.Response(503,json={'ack':'ok'}))) as c:
        assert worker.deliver(c,body) is None
    saved=list(tmp_path.glob('*.json'))
    assert len(saved)==1 and json.loads(saved[0].read_text())==body
    with httpx.Client(transport=httpx.MockTransport(lambda req: httpx.Response(200,json={'ack':'already_done'}))) as c:
        worker.flush_outbox(c)
    assert not list(tmp_path.glob('*.json'))


def test_worker_does_not_accept_unknown_ack(monkeypatch, tmp_path):
    import httpx
    monkeypatch.setenv('APP_URL','http://local.test');monkeypatch.setenv('WORKER_TOKEN','x'*40)
    from worker import box_worker as worker
    monkeypatch.setattr(worker,'OUTBOX',tmp_path);monkeypatch.setattr(worker.time,'sleep',lambda _:None)
    with httpx.Client(transport=httpx.MockTransport(lambda req:httpx.Response(200,json={'ack':'no_change'}))) as c:
        assert worker.deliver(c,{'id':9,'lease_id':'x','text':'{}'}) is None
    assert list(tmp_path.glob('*.json'))


def test_ai_reads_order_rows_on_later_workbook_sheets(tmp_path):
    import openpyxl
    from app.extract import _as_text
    wb=openpyxl.Workbook();wb.active.title='Cover';wb.active.append(['Contract summary'])
    ws=wb.create_sheet('Order detail');ws.append(['PO SECOND-SHEET']);ws.append(['Color','S','M']);ws.append(['Navy',10,20])
    path=tmp_path/'order.xlsx';wb.save(path)
    text=_as_text(path,'xlsx')
    assert 'SECOND-SHEET' in text and 'Navy' in text


def test_quantity_only_order_has_no_false_arithmetic_warning(client):
    body={'pos':[{'po_number':'QTYONLY','stated_total_qty':1200,'stated_total_value':3000,
                  'lines':[{'color':'Navy','sizes':{},'order_qty':1200,'rate':2.5,'net_value':3000}]}]}
    r=client.post('/api/orders',json=body).json()
    assert not any('sum of sizes' in w or 'rate x qty' in w for w in r['warnings'])


def test_queued_retry_updates_original_batch_result(client):
    bid='queued-recovery';old=db.add_file('pending.csv',status='queued_ai')
    result={'items':[{'file_id':old,'name':'pending.csv','status':'queued_ai','category':'csv'}]}
    result['summary']=main.summarize(result['items'],'pending.csv',False)
    db.save_batch(bid,'done','pending.csv',result)
    main._record_retry(old,{'file_id':old+1,'name':'pending.csv','status':'done','order_id':123,'category':'csv'})
    b=db.get_batch(bid)
    assert b['result']['items'][0]['status']=='done'
    assert b['result']['summary']['queued_ai']==0 and b['result']['summary']['extracted_ok']==1


def test_manual_edit_invalidates_ai_agreement_and_source_cache(client):
    oid=client.post('/api/orders',json={'pos':[{'po_number':'EDIT','lines':[{'color':'Navy','sizes':{'S':10}}]}]}).json()['id']
    fid=db.add_file('edited.csv',sha='edit-sha',status='done');db.update_file(fid,method='csv_ai_dual')
    with db.conn() as c:c.execute('UPDATE orders SET file_id=? WHERE id=?',(fid,oid))
    db.set_review(oid,'accepted',{'status':'accepted'})
    assert db.cached_result('edit-sha')
    assert client.put(f'/api/orders/{oid}',json={'pos':[{'po_number':'EDIT','lines':[{'color':'Navy','sizes':{'S':999}}]}]}).json()['ok']
    o=client.get(f'/api/orders/{oid}').json()
    assert o['extraction_status']!='accepted' and not db.cached_result('edit-sha')


def test_duplicate_po_blocks_cannot_hide_disagreement():
    a={'pos':[{'po_number':'P1','lines':[{'color':'Navy','sizes':{'S':1}}]}, {'po_number':'P1','lines':[{'color':'Navy','sizes':{'S':2}}]}]}
    b=json.loads(json.dumps(a));b['pos'][0]['lines'][0]['sizes']['S']=999
    assert consensus.compare(a,b)


def test_cover_sheet_does_not_skip_order_on_later_sheet(tmp_path, monkeypatch):
    import openpyxl
    from app import extract
    wb=openpyxl.Workbook();wb.active.append(['Contract summary'])
    ws=wb.create_sheet('Orders');ws.append(['Purchase Order PO12345']);ws.append(['Color','S','M']);ws.append(['Navy',10,20])
    p=tmp_path/'unknown.xlsx';wb.save(p)
    monkeypatch.setattr(extract,'ai_primary',lambda ctx:True)
    monkeypatch.setattr(extract,'extract_ai_primary',lambda *a:({'pos':[]},'test'))
    assert extract.extract_file(p,'xlsx',lambda *a:None)[1]=='test'


def test_size_aliases_are_summed_in_comparison():
    a={'pos':[{'po_number':'P1','lines':[{'color':'Navy','sizes':{'SM':10,'S':20}}]}]}
    b={'pos':[{'po_number':'P1','lines':[{'color':'Navy','sizes':{'S':20}}]}]}
    assert consensus.compare(a,b)


def test_quantity_without_sizes_must_exist_in_source():
    d={'pos':[{'po_number':'P1','lines':[{'color':'Navy','sizes':{},'order_qty':1200}]}]}
    assert guard.text_errors(d,'PO P1 Navy order quantity 500')


def test_text_groups_cover_every_page_without_exceeding_read_limit():
    from app.extract import _bounded_groups
    texts=['a'*30000,'b'*30000,'c'*12000,'d'*500]
    groups=_bounded_groups(texts,[(0,4,1)])
    assert [i for lo,hi,_ in groups for i in range(lo,hi)]==[0,1,2,3]
    assert all(len('\n\f'.join(texts[lo:hi]))<=60000 for lo,hi,_ in groups)


def test_merged_chunks_recheck_totals(monkeypatch):
    from app import extract
    monkeypatch.setattr(extract,'_need_llm',lambda *a:None)
    monkeypatch.setattr(extract,'_pdf_units',lambda *a:iter([{'text':'PO P1 Navy S 10 Total 10','page':'1','det':None},{'text':'PO P1 Red S 10 Total 10','page':'2','det':None}]))
    monkeypatch.setattr(llm,'read_page',lambda **kw:{'pos':[{'po_number':'P1','stated_total_qty':10,'lines':[{'color':'Navy' if 'Navy' in kw['text'] else 'Red','sizes':{'S':10}}]}]})
    ctx={};data,_=extract.extract_ai_primary(Path('unused.pdf'),'pdf',lambda *a:None,ctx)
    assert sum(sum(l['sizes'].values()) for l in data['pos'][0]['lines'])==20
    assert ctx['consensus']['status']=='needs_review'
    assert ctx['consensus']['arithmetic_errors']


def test_worker_progress_is_tied_to_active_lease(client, monkeypatch):
    token='p'*40;monkeypatch.setenv('WORKER_TOKEN',token)
    fid=db.add_file('progress-po.txt');jid=db.enqueue_job('prompt','model','text',None,owner='tester',file_id=fid,page='1-2',vote='A')
    job=db.claim_job('progress-worker','tester')
    while job and job['id']!=jid:
        db.finish_job(job['id'],job['lease_id'],result='{}');job=db.claim_job('progress-worker','tester')
    header={'Authorization':'Bearer '+token};body={'id':jid,'lease_id':job['lease_id'],'stage':'model_started'}
    assert client.post('/api/worker/progress',headers=header,json=body).json()['ok']
    snapshot=db.file_job_progress([fid])
    assert snapshot[fid]['jobs'][0]['stage']=='model_started' and snapshot[fid]['jobs'][0]['worker']=='progress-worker'
    assert not client.post('/api/worker/progress',headers=header,json={**body,'lease_id':'stale'}).json()['ok']
    db.finish_job(jid,job['lease_id'],result='{}')
    assert db.file_job_progress([fid])[fid]['done']==1


def test_progress_reports_real_worker_stage_and_counts(client):
    fid=db.add_file('stream.txt');jid=db.enqueue_job('p','model','text',None,owner='tester',file_id=fid,page='1',vote='A')
    main.PROG['stream-test']={'name':'stream.txt','started':time.time(),'files':[{'name':'stream.txt','file_id':fid,'step':'reading','reason':None,'order_id':None,'t0':time.time(),'t1':None,'order_doc':True}]}
    p=client.get('/api/progress',params={'bid':'stream-test'}).json()
    f=p['files'][0]
    assert f['worker_stage']=='queued' and f['reads']['total']==1 and f['reads']['done']==0
    db.cancel_job(jid)


def test_progress_separates_unprepared_files_and_cancelled_reads(client):
    fid=db.add_file('timeout.txt', status='queued_ai', reason='AI job timed out after 900s')
    jid=db.enqueue_job('p','m','text',None,file_id=fid)
    db.cancel_job(jid)
    queued_fid=db.add_file('queued.txt')
    db.enqueue_job('p','m','text',None,file_id=queued_fid)
    def entry(name, file_id=None, step='queued'):
        return dict(name=name,file_id=file_id,step=step,reason=None,order_id=None,t0=None,t1=None,order_doc=True)
    main.PROG['honest-progress']={'name':'input.zip','started':time.time(),'files':[
        entry('timeout.txt',fid),entry('queued.txt',queued_fid),*[entry(f'pending-{i}.txt') for i in range(45)]]}
    p=main.progress('honest-progress')
    assert p['files'][0]['worker_stage']=='retry_pending'
    assert p['files'][2]['worker_stage']=='pending'
    assert p['reads_failed']==1 and p['reads_done']==0 and p['reads_total']==2
    assert p['reads_queued']==1 and p['files_pending']==45 and p['files_retrying']==1
    assert p['files'][0]['reads']['jobs'][0]['stage']=='cancelled'


def test_cancelled_batch_never_shows_active_processing(client):
    db.save_batch('cancelled-progress','cancelled','input.zip',{'error':'Queue cleared once'})
    p=main.progress('cancelled-progress')
    assert p is not None and not p['active']
    assert p['batch_status']=='cancelled' and p['error']=='Queue cleared once'


def test_empty_worker_response_has_no_body():
    db.init()
    response=main.worker_next(owner='wrong-account')
    assert response.status_code==204 and response.body==b''


def test_worker_milestones_keep_receipt_and_start_times():
    db.init();fid=db.add_file('milestones.txt');jid=db.enqueue_job('p','m','text',None,owner='milestone-account',file_id=fid)
    job=db.claim_job('grok-test','milestone-account')
    assert db.update_job_progress(jid,job['lease_id'],'model_started')
    assert db.finish_job(jid,job['lease_id'],result='{}')=='ok'
    record=db.file_job_progress([fid])[fid]['jobs'][0]
    assert record['received_at'] and record['started_at'] and record['returned_at']
    assert record['received_at'] <= record['started_at'] <= record['returned_at']


def test_empty_po_blocks_are_not_accepted(monkeypatch):
    from app import extract
    monkeypatch.setattr(extract,'_need_llm',lambda *a:None)
    monkeypatch.setattr(extract,'_pdf_units',lambda *a:iter([{'text':'PO P1','page':'1','det':None}]))
    monkeypatch.setattr(llm,'read_page',lambda **kw:{'pos':[{'po_number':'P1','lines':[]}]})
    ctx={};extract.extract_ai_primary(Path('unused.pdf'),'pdf',lambda *a:None,ctx)
    assert ctx['consensus']['status']=='needs_review'


def test_waiting_upload_is_not_shown_as_finished(client):
    bid='waiting-upload';db.save_batch(bid,'processing','waiting.zip')
    main.PROG[bid]={'name':'waiting.zip','started':time.time(),'files':[]}
    p=client.get('/api/progress',params={'bid':bid}).json()
    assert p['active'] and p['phase']=='waiting'
    db.save_batch(bid,'failed','waiting.zip',{'error':'QA cleanup'})


def test_retry_progress_binds_new_jobs(client):
    old=db.add_file('retry-progress.txt',status='queued_ai');new=db.add_file('retry-progress.txt')
    bid='bind-retry';db.save_batch(bid,'done','retry-progress.txt',{'items':[{'file_id':old,'name':'retry-progress.txt','status':'queued_ai'}]})
    main.PROG[bid]={'name':'retry-progress.txt','started':time.time(),'files':[{'name':'retry-progress.txt','file_id':old,'step':'queued','reason':None,'order_id':None,'t0':None,'t1':None,'order_doc':True}]}
    main._bind_retry(old,new)
    jid=db.enqueue_job('p','m','text',None,owner='retry-account',file_id=new)
    job=db.claim_job('retry-worker','retry-account');db.update_job_progress(jid,job['lease_id'],'model_started')
    p=main.progress(bid)
    assert p['files'][0]['worker_stage']=='model_started'
    db.finish_job(jid,job['lease_id'],result='{}')
    main._record_retry(old,{'file_id':new,'name':'retry-progress.txt','status':'done','order_id':123,'category':'other'})
    p=main.progress(bid)
    assert p['reads_done']==1 and p['files'][0]['reads']['jobs'][0]['id']==jid


def test_slow_event_callback_cannot_accept_after_timeout(monkeypatch, tmp_path):
    from app import opencode_provider as oc
    exe=tmp_path/'fake-opencode'
    exe.write_text('#!/usr/bin/env python3\nimport json,sys\nevents=[{"type":"text","part":{"text":"x"}}]*12+[{"type":"step_finish","part":{"reason":"stop"}}]\nsys.stdout.write("\\n".join(json.dumps(e) for e in events)+"\\n");sys.stdout.flush()\n')
    exe.chmod(0o755);monkeypatch.setattr(oc,'binary',lambda:str(exe))
    with pytest.raises(TimeoutError):oc.run('p',timeout=.15,on_event=lambda event:time.sleep(.04))


def test_worker_shutdown_cancels_its_model_process(monkeypatch, tmp_path):
    import threading
    from app import opencode_provider as oc
    exe=tmp_path/'fake-opencode';exe.write_text('#!/usr/bin/env python3\nimport time\ntime.sleep(30)\n');exe.chmod(0o755)
    monkeypatch.setattr(oc,'binary',lambda:str(exe));stop=threading.Event();stop.set()
    start=time.monotonic()
    with pytest.raises(InterruptedError):oc.run('p',timeout=30,stop_event=stop)
    assert time.monotonic()-start < 2


def test_dropdown_choices_use_saved_orders_and_keep_unknown_values(client):
    body={'buyer_name':'Dropdown buyer','brand':'Real brand','factory':'Real factory','department':'Real department','sizes':['S','M'],
          'colors':['Navy'],'pos':[{'po_number':'SELECT-PO','destination':'Real destination','lines':[{'color':'Navy','sizes':{'S':10,'M':20}}]}]}
    oid=client.post('/api/orders',json=body).json()['id']
    options=db.order_options()
    for key,value in [('buyer_name','Dropdown buyer'),('brand','Real brand'),('factory','Real factory'),('department','Real department'),('destination','Real destination'),('colors','Navy')]:
        assert value in options[key]
    assert 'S, M' in options['sizes']
    page=client.get('/orders/new').text
    assert 'Real brand' in page and 'Real factory' in page
    client.delete(f'/api/orders/{oid}')
